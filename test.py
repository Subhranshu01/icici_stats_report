import os
import requests
import pandas as pd
from datetime import datetime, timedelta
from openpyxl import load_workbook
import smtplib
import pytz
from email.message import EmailMessage
from openpyxl.styles import Alignment, PatternFill

india_tz = pytz.timezone("Asia/Kolkata")
now_ist = datetime.now(india_tz)
FILE_NAME = "Product Category wise Internal APIs Performance Report.xlsx"

def apply_formatting(file_path):
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    wb = load_workbook(file_path)

    for sheet in wb.worksheets:
        header_row = next(sheet.iter_rows(min_row=1, max_row=1))
        col_map = {cell.value: cell.column for cell in header_row}

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.value is not None:
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

            cell_p90 = row[col_map.get("p90") - 1]
            if isinstance(cell_p90.value, (int, float)) and cell_p90.value > 3:
                cell_p90.fill = red_fill

            cell_p95 = row[col_map.get("p95") - 1]
            if isinstance(cell_p95.value, (int, float)) and cell_p95.value > 3:
                cell_p95.fill = red_fill

            cell_fr = row[col_map.get("failure rate") - 1]
            if isinstance(cell_fr.value, (int, float)) and cell_fr.value > 10:
                cell_fr.fill = red_fill

        for col in sheet.columns:
            max_length = max((len(str(cell.value)) if cell.value else 0) for cell in col)
            adjusted_width = max_length + 2
            col_letter = col[0].column_letter
            sheet.column_dimensions[col_letter].width = adjusted_width

    wb.save(file_path)
    print("🎨 Formatting + 🔴 highlights applied to all sheets.")

def fetch_and_store_metrics(controller_name, url, api_token):
    yesterday_str = (now_ist - timedelta(days=1)).strftime("%Y-%m-%d")
    sheet_name = controller_name

    headers = {
        "Authorization": f"Api-Token {api_token}",
        "accept": "application/json"
    }
    response = requests.get(url, headers=headers)
    print(f"📡 {controller_name}: Status Code {response.status_code}")
    if response.status_code != 200:
        print("❌ Error:", response.text)
        return

    json_data = response.json()
    result = json_data.get("result", [])
    data_dict = {}

    for metric in result:
        metric_id = metric.get("metricId", "")
        print(f"\n🔍 Processing Metric ID: {metric_id}")
        for entry in metric.get("data", []):
            dimension_map = entry.get("dimensionMap", {})
            method_name = (
                dimension_map.get("dt.entity.service_method.name") or
                dimension_map.get("Dimension") or
                "unknown_method"
            )

            value = entry.get("values", [0])[0]
            print(f"📛 Method: {method_name} | Value: {value}")
            if method_name not in data_dict:
                data_dict.setdefault(method_name, {})

            if "count.total" in metric_id or "total_count" in metric_id:
                data_dict[method_name]["total_hits"] = int(value)
                print(f"✅ total_hits assigned for {method_name}")
            elif "errors.server.count" in metric_id or "failed_req" in metric_id:
                data_dict[method_name]["failure_count"] = int(value)
                print(f"✅ failure_count assigned for {method_name}")
            elif ":avg" in metric_id or "avg_responsetime" in metric_id:
                data_dict[method_name]["avg"] = round(value / 1_000_000, 2)
                print(f"✅ average assigned for {method_name}")
            elif "percentile(90.0)" in metric_id:
                data_dict[method_name]["p90"] = round(value / 1_000_000, 2)
            elif "percentile(95.0)" in metric_id:
                data_dict[method_name]["p95"] = round(value / 1_000_000, 2)
            elif "percentile(99.0)" in metric_id:
                data_dict[method_name]["p99"] = round(value / 1_000_000, 2)
            elif "errors.server.rate" in metric_id:
                data_dict[method_name]["failure_rate"] = round(value, 2)

    records = []
    for method, values in data_dict.items():
        if "total_hits" not in values:
            print(f"⚠️ total_hits missing for {method}, defaulting to 0")
        if "failure_count" not in values:
            print(f"⚠️ failure_count missing for {method}, defaulting to 0")
        if "avg" not in values:
            print(f"⚠️ average missing for {method}, defaulting to 0")
        
        records.append({
            "Date": yesterday_str,
            "request": method,
            "total hits": values.get("total_hits", 0),
            "failure count": values.get("failure_count", 0),
            "average": values.get("avg", 0),
            "p90": values.get("p90", 0),
            "p95": values.get("p95", 0),
            "p99": values.get("p99", 0),
            "failure rate": values.get("failure_rate", 0)
        })

    df = pd.DataFrame(records)
    print("📊 Final DataFrame:")
    print(df.head())
    update_workbook(sheet_name, df)

def update_workbook(sheet_name, df_new):
    all_sheets = {}

    if os.path.exists(FILE_NAME):
        with pd.ExcelFile(FILE_NAME, engine="openpyxl") as xls:
            for name in xls.sheet_names:
                try:
                    all_sheets[name] = pd.read_excel(xls, sheet_name=name)
                except Exception:
                    all_sheets[name] = pd.DataFrame()
        print(f"📄 Existing sheets loaded: {list(all_sheets.keys())}")
    else:
        print("🆕 No existing workbook found. Creating new one.")

    existing_df = all_sheets.get(sheet_name, pd.DataFrame())
    combined_df = pd.concat([df_new, existing_df], ignore_index=True)
    all_sheets[sheet_name] = combined_df

    with pd.ExcelWriter(FILE_NAME, engine="openpyxl", mode="w") as writer:
        for name, df in all_sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
    print(f"✅ Sheet '{sheet_name}' updated with new data.")

def send_email_report():
    EMAIL_USER = os.environ["EMAIL_USER"]
    EMAIL_PASS = os.environ["EMAIL_PASS"]
    SMTP_SERVER = os.environ["EMAIL_HOST"]
    SMTP_PORT = int(os.environ["EMAIL_PORT"])
    TO_EMAIL = os.environ["EMAIL_TO"]

    msg = EmailMessage()
    msg["Subject"] = "📊 Dynatrace Metrics Report"
    msg["From"] = EMAIL_USER
    msg["To"] = TO_EMAIL
    msg.set_content("Hi,\n\nAttached is the updated Product Category wise Internal APIs Performance Report.\n\nRegards,\nSubhranshu")

    with open(FILE_NAME, "rb") as f:
        msg.add_attachment(
            f.read(),
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=FILE_NAME
        )

    try:
        with smtplib.SMTP(SMTP_SERVER, SMTP_PORT) as smtp:
            smtp.starttls()
            smtp.login(EMAIL_USER, EMAIL_PASS)
            smtp.send_message(msg)
        print("📤 Email sent successfully.")
    except Exception as e:
        print("❌ Email sending failed:", e)

if __name__ == "__main__":
    API_TOKEN = os.environ["API_TOKEN"]
    LoginController_url = os.environ["LOGINCONTROLLER_URL"]
    MotorInsurance_url = os.environ["MOTORINSURANCE_URL"]
    CreditTrack_url = os.environ["CREDITTRACK_URL"]
    Digigold_url = os.environ["DIGIGOLD_URL"]
    HealthInsurance_url = os.environ["HEALTHINSURANCE_URL"]
    LrRewards_url = os.environ["LRREWARDS_URL"]
    PersonalLoan_url = os.environ["PERSONALLOAN_URL"]
    MutualFund_url = os.environ["MUTUALFUND_URL"]
    FixedDeposit_url = os.environ["FIXEDDEPOSIT_URL"]
    BusinessLoanController_url = os.environ["BLOAN_URL"]
    TermInsuranceBuyController_url = os.environ["TERM_URL"]
    Stocks_url = os.environ["STOCKS_URL"]
    GoldLoanController_url = os.environ["GOLD_URL"]
    HomeLoan_url = os.environ["HOMELOAN_URL"]
    PortfolioTrack_url = os.environ["PORTFOLIO_URL"]
    SpendTrack_url = os.environ["SPENDTRACK_URL"]

    fetch_and_store_metrics("LoginController", LoginController_url, API_TOKEN)
    fetch_and_store_metrics("MotorInsurance", MotorInsurance_url, API_TOKEN)
    fetch_and_store_metrics("CreditTrack", CreditTrack_url, API_TOKEN)
    fetch_and_store_metrics("DigiGold & Silver", Digigold_url, API_TOKEN)
    fetch_and_store_metrics("HealthInsurance", HealthInsurance_url, API_TOKEN)
    fetch_and_store_metrics("LR Rewards", LrRewards_url, API_TOKEN)
    fetch_and_store_metrics("Personal Loan", PersonalLoan_url, API_TOKEN)
    fetch_and_store_metrics("MutualFund", MutualFund_url, API_TOKEN)
    fetch_and_store_metrics("Fixed Deposit", FixedDeposit_url, API_TOKEN)
    fetch_and_store_metrics("BusinessLoanController", BusinessLoanController_url, API_TOKEN)
    fetch_and_store_metrics("TermInsuranceBuyController(LI)", TermInsuranceBuyController_url, API_TOKEN)
    fetch_and_store_metrics("Stocks", Stocks_url, API_TOKEN)
    fetch_and_store_metrics("GoldLoanController", GoldLoanController_url, API_TOKEN)
    fetch_and_store_metrics("Home Loan", HomeLoan_url, API_TOKEN)
    fetch_and_store_metrics("PortfolioTrack", PortfolioTrack_url, API_TOKEN)
    fetch_and_store_metrics("SpendTrack", SpendTrack_url, API_TOKEN)
    
    apply_formatting(FILE_NAME)
    send_email_report()






